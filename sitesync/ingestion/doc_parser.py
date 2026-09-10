"""
SiteSync AI — Multi-Modal Document Parser

Priority order:
  1. Azure Document Intelligence (if configured) — best OCR for scanned docs
  2. pdfplumber (PDF text extraction fallback)
  3. openpyxl (Excel/XLSX)
  4. Plain text passthrough

Returns a unified ParsedDocument dataclass.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from rich.console import Console

from sitesync.config import settings

console = Console()


@dataclass
class ParsedDocument:
    """Unified output from any document parser."""
    raw_text: str
    tables: list[list[list[str]]] = field(default_factory=list)  # list of tables → rows → cells
    source_type: str = "text"  # text | pdf | xlsx | image | scanned_doc
    evidence_type: str = "text"  # mapped to schema
    filename: str = ""
    page_count: int = 0
    parser_used: str = "passthrough"
    error: Optional[str] = None


def _evidence_from_source(source_type: str) -> str:
    mapping = {
        "pdf": "scanned_doc",
        "image": "photo",
        "xlsx": "spreadsheet",
        "csv": "spreadsheet",
        "text": "text",
    }
    return mapping.get(source_type, "text")


def _parse_with_azure_di(file_bytes: bytes, content_type: str) -> ParsedDocument:
    """Parse using Azure Document Intelligence Layout model."""
    from azure.ai.documentintelligence import DocumentIntelligenceClient
    from azure.ai.documentintelligence.models import AnalyzeDocumentRequest, ContentFormat
    from azure.core.credentials import AzureKeyCredential

    client = DocumentIntelligenceClient(
        endpoint=settings.azure_document_intelligence_endpoint,
        credential=AzureKeyCredential(settings.azure_document_intelligence_key),
    )

    poller = client.begin_analyze_document(
        "prebuilt-layout",
        analyze_request=file_bytes,
        content_type=content_type,
        output_content_format=ContentFormat.TEXT,
    )
    result = poller.result()

    raw_text = result.content or ""
    tables: list[list[list[str]]] = []

    if result.tables:
        for table in result.tables:
            row_count = table.row_count
            col_count = table.column_count
            grid = [[""] * col_count for _ in range(row_count)]
            for cell in table.cells:
                grid[cell.row_index][cell.column_index] = cell.content or ""
            tables.append(grid)

    return ParsedDocument(
        raw_text=raw_text,
        tables=tables,
        source_type="pdf",
        evidence_type="scanned_doc",
        page_count=len(result.pages) if result.pages else 0,
        parser_used="azure_document_intelligence",
    )


def _parse_pdf_fallback(file_bytes: bytes) -> ParsedDocument:
    """Parse PDF using pdfplumber."""
    import pdfplumber

    text_parts = []
    tables = []

    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        page_count = len(pdf.pages)
        for page in pdf.pages:
            text_parts.append(page.extract_text() or "")
            page_tables = page.extract_tables()
            if page_tables:
                tables.extend(page_tables)

    return ParsedDocument(
        raw_text="\n\n".join(filter(None, text_parts)),
        tables=tables,
        source_type="pdf",
        evidence_type="scanned_doc",
        page_count=page_count,
        parser_used="pdfplumber",
    )


def _parse_xlsx(file_bytes: bytes, filename: str = "") -> ParsedDocument:
    """Parse Excel using openpyxl — extracts all sheets as tables + text."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    all_text_parts = []
    all_tables = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows_data = []
        text_rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            # Skip completely empty rows
            if any(c.strip() for c in cells):
                rows_data.append(cells)
                text_rows.append("\t".join(cells))

        if rows_data:
            all_tables.append(rows_data)
            all_text_parts.append(f"[Sheet: {sheet_name}]\n" + "\n".join(text_rows))

    raw_text = "\n\n".join(all_text_parts)

    return ParsedDocument(
        raw_text=raw_text,
        tables=all_tables,
        source_type="xlsx",
        evidence_type="spreadsheet",
        filename=filename,
        parser_used="openpyxl",
    )


def parse_file(
    file_bytes: bytes,
    filename: str,
    force_parser: Optional[str] = None,
) -> ParsedDocument:
    """
    Auto-detect file type and parse with the best available parser.

    Args:
        file_bytes: Raw file content.
        filename: Original filename (used for extension detection).
        force_parser: Force a specific parser ('azure', 'pdfplumber', 'openpyxl').
    """
    suffix = Path(filename).suffix.lower()

    # ── Excel ─────────────────────────────────────────────────────────────────
    if suffix in (".xlsx", ".xls", ".xlsm") or force_parser == "openpyxl":
        doc = _parse_xlsx(file_bytes, filename=filename)
        doc.filename = filename
        return doc

    # ── PDF or Image → try Azure DI first ─────────────────────────────────────
    if suffix in (".pdf", ".png", ".jpg", ".jpeg", ".tiff", ".bmp") or force_parser in ("azure", "pdfplumber"):
        content_type = "application/pdf" if suffix == ".pdf" else "image/jpeg"

        if settings.has_azure_di and force_parser != "pdfplumber":
            try:
                doc = _parse_with_azure_di(file_bytes, content_type)
                doc.filename = filename
                return doc
            except Exception as e:
                console.print(f"[yellow]⚠ Azure DI failed ({e!s:.60}), falling back to pdfplumber[/yellow]")

        if suffix == ".pdf":
            doc = _parse_pdf_fallback(file_bytes)
            doc.filename = filename
            return doc
        else:
            return ParsedDocument(
                raw_text="",
                filename=filename,
                source_type="image",
                evidence_type="photo",
                error="Image parsing requires Azure DI (not configured)",
                parser_used="none",
            )

    # ── Plain text ─────────────────────────────────────────────────────────────
    try:
        text = file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        text = file_bytes.decode("latin-1", errors="replace")

    return ParsedDocument(
        raw_text=text,
        filename=filename,
        source_type="text",
        evidence_type="text",
        parser_used="passthrough",
    )


def parse_text(text: str) -> ParsedDocument:
    """Parse a plain text string directly (for the chat Time Agent)."""
    return ParsedDocument(
        raw_text=text,
        source_type="text",
        evidence_type="text",
        parser_used="passthrough",
    )
